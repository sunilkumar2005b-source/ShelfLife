from pathlib import Path
from pptx import Presentation

target = Path(r"C:\Users\SUNILKUMAR\Downloads\ShelfLife.pptx")
prs = Presentation(target)

print(f"Presentation Path: {target}")
print(f"Total Slides: {len(prs.slides)}")
print(f"Widescreen Dimensions: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches")

for idx, slide in enumerate(prs.slides):
    texts = []
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text.strip():
            texts.append(s.text_frame.text.strip().replace("\n", " -- "))
    images = [s for s in slide.shapes if s.shape_type == 13]
    tables = [s for s in slide.shapes if s.has_table]
    main_title = texts[1] if len(texts) > 1 and "SWE3004" in texts[0] else (texts[0] if texts else "NO TEXT")
    print(f"Slide {idx+1:02d} | Shapes: {len(slide.shapes):2d} | Img: {len(images)} | Tbl: {len(tables)} | {main_title[:65]}")

print("\nVerification Completed Successfully!")
