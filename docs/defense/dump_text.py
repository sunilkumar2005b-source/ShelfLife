import sys
from pptx import Presentation

prs = Presentation(r'C:\Users\SUNILKUMAR\Downloads\ShelfLife.pptx')
with open(r'C:\Users\SUNILKUMAR\OneDrive\Desktop\ShelfLife\docs\defense\existing_slides_text.txt', 'w', encoding='utf-8') as f:
    for idx, slide in enumerate(prs.slides):
        f.write(f"\n==================== SLIDE {idx+1} ====================\n")
        f.write(f"Layout: {slide.slide_layout.name}\n")
        for s in slide.shapes:
            if s.has_text_frame:
                f.write(f"--- Shape [{s.name}] ({s.shape_type}) pos=({s.left.inches:.2f}, {s.top.inches:.2f}) dim=({s.width.inches:.2f}x{s.height.inches:.2f}) ---\n")
                for p in s.text_frame.paragraphs:
                    txt = p.text.strip()
                    if txt:
                        f.write(f"    {txt}\n")
            elif s.has_table:
                f.write(f"--- Table [{s.name}] pos=({s.left.inches:.2f}, {s.top.inches:.2f}) ---\n")
                for row in s.table.rows:
                    row_txt = [c.text.strip() for c in row.cells]
                    f.write(f"    {row_txt}\n")
            elif s.shape_type == 13: # PICTURE
                f.write(f"--- Picture [{s.name}] pos=({s.left.inches:.2f}, {s.top.inches:.2f}) dim=({s.width.inches:.2f}x{s.height.inches:.2f}) ---\n")

print("Dumped existing slides text successfully.")
