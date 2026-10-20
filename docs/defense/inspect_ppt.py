import sys
from pptx import Presentation

prs = Presentation(r'C:\Users\SUNILKUMAR\Downloads\ShelfLife.pptx')
print(f'Slide width: {prs.slide_width.inches:.2f} in, height: {prs.slide_height.inches:.2f} in')
print(f'Total slides: {len(prs.slides)}')
for i in range(len(prs.slides)):
    sl = prs.slides[i]
    print(f"\n--- Slide {i+1} (Layout: {sl.slide_layout.name}) ---")
    for s in sl.shapes:
        txt = s.text_frame.text.replace('\n', ' ')[:60] if s.has_text_frame else ''
        print(f"  Shape [{s.name}] type={s.shape_type} pos=({s.left.inches:.2f}, {s.top.inches:.2f}) dim=({s.width.inches:.2f}x{s.height.inches:.2f}) text='{txt}'")
