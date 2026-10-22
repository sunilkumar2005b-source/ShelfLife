# ==============================================================================
# ShelfLife - Master Expanded Dissertation Generator (Target: 75-85 Pages)
# Implements full academic depth, detailed tables, code listings, and figures
# ==============================================================================

import os
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCS_DIR = Path(r"C:\Users\SUNILKUMAR\OneDrive\Desktop\ShelfLife\docs\defense")
TARGET_DOCX = DOCS_DIR / "SWE3004_Dissertation_ShelfLife.docx"
FINAL_DOCX = DOCS_DIR / "SWE3004_Final_Project_Report_ShelfLife.docx"

# Assets
IMG_LOGO = DOCS_DIR / "extracted_images" / "image1.jpg"
IMG_ARCH = DOCS_DIR / "system_architecture_diagram.png"
IMG_CLASS = DOCS_DIR / "uml_class_diagram.png"
IMG_USECASE = DOCS_DIR / "uml_use_case_diagram.png"
IMG_WORKFLOW = DOCS_DIR / "activity_workflow_diagram.png"
IMG_GANTT = DOCS_DIR / "project_schedule_gantt.png"
IMG_DASH = DOCS_DIR / "extracted_images" / "image5.png"
IMG_BATCH = DOCS_DIR / "extracted_images" / "image6.png"
IMG_AI = DOCS_DIR / "extracted_images" / "image7.png"
IMG_POS = DOCS_DIR / "extracted_images" / "image8.png"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="999999", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_field(p, field_text):
    run = p.add_run()
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> %s </w:instrText>' % (nsdecls('w'), field_text))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_custom_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(28)
        p.paragraph_format.space_after = Pt(14)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.bold = True
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.bold = True
    elif level == 3:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.bold = True
    return p

def add_body_p(doc, text, space_after=6, bold=False, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic
    return p

def add_bullet_p(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.bold = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return p

def add_numbered_p(doc, num_str, bold_prefix, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    r_num = p.add_run(num_str + " ")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(12)
    r_num.bold = True
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.bold = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return p

def add_caption(doc, text, is_table=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = is_table
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.bold = True
    return p

def add_figure(doc, img_path, caption_text, width_inches=6.0):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(str(img_path), width=Inches(width_inches))
        add_caption(doc, caption_text, is_table=False)
    else:
        print(f"Warning: Image {img_path} not found.")

def add_styled_table(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="777777", sz="4", val="single")

    # Header row
    hdr_row = table.rows[0]
    hdr_tr = hdr_row._tr.get_or_add_trPr()
    hdr_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    for idx, text in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "E6E8EC")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(str(text))
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
        run.bold = True

    # Data rows
    for r_idx, row_values in enumerate(rows_data):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            val_str = str(val)
            if c_idx == 0 and len(val_str) <= 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif c_idx == len(row_values) - 1 and len(val_str) <= 14:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(val_str)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)

    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def add_code_block(doc, title, code_text):
    if title:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(10)
        p_t.paragraph_format.space_after = Pt(2)
        r_t = p_t.add_run(title)
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(11)
        r_t.bold = True
    
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="CBD5E1", sz="6", val="single")
    cell = tbl.rows[0].cells[0]
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    cell.width = Inches(6.5)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.0)
    run.font.color.rgb = RGBColor(30, 41, 59)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

print("Base setup loaded.")
